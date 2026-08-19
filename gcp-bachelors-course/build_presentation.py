#!/usr/bin/env python3
"""Build GCP bachelor CS classroom PowerPoint from course materials."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).parent / "GCP-Bachelor-CS-Classroom.pptx"

# Google-inspired palette (teaching use)
BLUE = RGBColor(0x1A, 0x73, 0xE8)
DARK = RGBColor(0x20, 0x21, 0x24)
GREEN = RGBColor(0x1E, 0x8E, 0x3E)
ORANGE = RGBColor(0xE3, 0x74, 0x00)
GREY = RGBColor(0x5F, 0x63, 0x68)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF8, 0xF9, 0xFA)


def set_slide_bg(slide, color=LIGHT_BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title_slide(prs, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, Inches(0.15))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()

    tx = slide.shapes.add_textbox(Inches(0.6), Inches(2.0), Inches(8.8), Inches(1.5))
    p = tx.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = DARK

    if subtitle:
        st = slide.shapes.add_textbox(Inches(0.6), Inches(3.6), Inches(8.8), Inches(1.2))
        sp = st.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(20)
        sp.font.color.rgb = GREY


def add_section_slide(prs, week, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BLUE)
    tx = slide.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(8.4), Inches(2))
    tf = tx.text_frame
    p = tf.paragraphs[0]
    p.text = week
    p.font.size = Pt(22)
    p.font.color.rgb = WHITE
    p2 = tf.add_paragraph()
    p2.text = title
    p2.font.size = Pt(36)
    p2.font.bold = True
    p2.font.color.rgb = WHITE


def add_bullet_slide(prs, title, bullets, notes=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(9), Inches(0.8))
    tp = title_box.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(28)
    tp.font.bold = True
    tp.font.color.rgb = BLUE

    body = slide.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(8.6), Inches(5.8))
    tf = body.text_frame
    tf.word_wrap = True
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(18)
        p.font.color.rgb = DARK
        p.space_after = Pt(8)

    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def add_two_column_slide(prs, title, left_title, left_items, right_title, right_items, notes=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(9), Inches(0.7))
    title_box.text_frame.paragraphs[0].text = title
    title_box.text_frame.paragraphs[0].font.size = Pt(28)
    title_box.text_frame.paragraphs[0].font.bold = True
    title_box.text_frame.paragraphs[0].font.color.rgb = BLUE

    for col, (ctitle, items, x) in enumerate(
        [(left_title, left_items, 0.5), (right_title, right_items, 5.0)]
    ):
        hdr = slide.shapes.add_textbox(Inches(x), Inches(1.1), Inches(4.3), Inches(0.5))
        hp = hdr.text_frame.paragraphs[0]
        hp.text = ctitle
        hp.font.size = Pt(20)
        hp.font.bold = True
        hp.font.color.rgb = GREEN if col == 0 else ORANGE

        box = slide.shapes.add_textbox(Inches(x), Inches(1.65), Inches(4.3), Inches(5.2))
        tf = box.text_frame
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f"• {item}"
            p.font.size = Pt(16)
            p.font.color.rgb = DARK
            p.space_after = Pt(6)

    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def add_table_slide(prs, title, headers, rows, notes=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(9), Inches(0.7))
    title_box.text_frame.paragraphs[0].text = title
    title_box.text_frame.paragraphs[0].font.size = Pt(26)
    title_box.text_frame.paragraphs[0].font.bold = True
    title_box.text_frame.paragraphs[0].font.color.rgb = BLUE

    cols, rcount = len(headers), len(rows) + 1
    table = slide.shapes.add_table(rcount, cols, Inches(0.5), Inches(1.2), Inches(9), Inches(0.4 * rcount)).table

    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(14)
            p.font.color.rgb = WHITE
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE

    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(13)
                p.font.color.rgb = DARK

    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def add_diagram_slide(prs, title, lines, notes=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(9), Inches(0.7))
    title_box.text_frame.paragraphs[0].text = title
    title_box.text_frame.paragraphs[0].font.size = Pt(26)
    title_box.text_frame.paragraphs[0].font.bold = True
    title_box.text_frame.paragraphs[0].font.color.rgb = BLUE

    box = slide.shapes.add_shape(1, Inches(0.7), Inches(1.2), Inches(8.6), Inches(5.5))
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = BLUE

    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Consolas"
        p.font.size = Pt(15)
        p.font.color.rgb = DARK
        p.alignment = PP_ALIGN.CENTER

    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def build():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # --- Title ---
    add_title_slide(
        prs,
        "Google Cloud Platform",
        "Bachelor of Computer Science — 12-Week Course\nCompute · Network · Data · Real-Time Systems",
    )

    add_bullet_slide(
        prs,
        "Course at a glance",
        [
            "Level: B.Sc. / B.Tech. CS (Year 2–3)",
            "36 lecture hours + 24 lab hours",
            "Prerequisites: OS, networks, databases, Linux, one programming language",
            "Goal: design, deploy, observe, and cost a GCP application",
            "Capstone: real-time path (Pub/Sub, Dataflow, or Firestore listeners)",
        ],
        notes="Open with: we teach decisions, not product catalogs.",
    )

    add_bullet_slide(
        prs,
        "Learning outcomes",
        [
            "Explain IaaS / PaaS / SaaS / FaaS and map to GCP products",
            "Apply IAM least privilege across projects and service accounts",
            "Choose compute: GCE, GKE, Cloud Run, Cloud Functions",
            "Design VPC, load balancing, CDN, and Cloud Armor",
            "Select storage & databases by consistency, latency, and cost",
            "Build event-driven pipelines: Pub/Sub → Dataflow → BigQuery",
            "Instrument apps with Logging, Monitoring, Trace, SLOs",
            "Critique real-time architectures (gaming, IoT, fraud, media)",
        ],
    )

    # --- Week 1 ---
    add_section_slide(prs, "Week 1", "Cloud Computing & GCP Mental Model")

    add_two_column_slide(
        prs,
        "Why cloud? (CS view)",
        "Problems in a data center",
        [
            "Capacity planning before demand",
            "Hardware & zone failures",
            "Multi-tenant isolation",
            "Slow manual provisioning",
        ],
        "NIST cloud characteristics",
        [
            "On-demand self-service (APIs)",
            "Broad network access",
            "Resource pooling",
            "Rapid elasticity",
            "Measured service (billing)",
        ],
        notes="Ask: what is still YOUR responsibility after you move to cloud?",
    )

    add_table_slide(
        prs,
        "Service models → GCP",
        ["Model", "You manage", "GCP examples"],
        [
            ["IaaS", "OS, runtime, app", "Compute Engine, Persistent Disk"],
            ["CaaS", "Container image", "GKE, Cloud Run"],
            ["PaaS", "Application code", "App Engine, Cloud Run"],
            ["FaaS", "Function + trigger", "Cloud Functions"],
            ["SaaS", "Nothing (consumer)", "Gmail (not what we build)"],
        ],
    )

    add_diagram_slide(
        prs,
        "Regions, zones & edge",
        [
            "  [ User ]",
            "      ↓",
            "  Google Edge (CDN / Global LB)",
            "      ↓  (Google backbone)",
            "  Region: asia-south1",
            "    ├── Zone a   (failure domain)",
            "    ├── Zone b",
            "    └── Zone c",
            "",
            "Rule: put compute next to data",
        ],
        notes="Zone = failure domain. Region = latency & data residency.",
    )

    add_bullet_slide(
        prs,
        "Resource hierarchy & shared responsibility",
        [
            "Organization → Folders → Projects → Resources",
            "IAM inherits downward — grant at the highest correct node",
            "Google: physical DC, hypervisor, managed DB patches",
            "You: IAM policy, app authZ, GCE OS patches, data classification",
            "SLA (product) ≠ SLO (your user promise) — design for both",
        ],
    )

    # --- Week 2 ---
    add_section_slide(prs, "Week 2", "IAM, Identity & Billing")

    add_table_slide(
        prs,
        "IAM principals & roles",
        ["Principal", "Use case"],
        [
            ["Google Account", "Humans (faculty, students)"],
            ["Service Account", "Workloads (API, Cloud Run, GCE)"],
            ["Google Group", "Teams"],
            ["Workload Identity Federation", "CI/CD without JSON keys"],
        ],
        notes="Never download JSON keys for new designs. Use attached SAs.",
    )

    add_bullet_slide(
        prs,
        "IAM policy & FinOps basics",
        [
            "Policy = allow { member, role } on resource",
            "Prefer predefined/custom roles over Owner/Editor/Viewer",
            "IAM Conditions: time, IP, resource tags",
            "Billing: SKUs, labels, budgets + Pub/Sub alerts",
            "Surprise costs: egress, idle GKE, BigQuery SELECT *",
            "Quotas protect you — lab failures are often quota, not IAM",
        ],
    )

    # --- Week 3 ---
    add_section_slide(prs, "Week 3", "Compute Engine & MIGs")

    add_bullet_slide(
        prs,
        "GCE essentials",
        [
            "VM = machine type + boot disk + NIC + service account",
            "Families: general (E2/N2), compute (C2), memory (M3), GPU (A2/G2)",
            "Disks: pd-balanced, pd-ssd, Local SSD (ephemeral, high IOPS)",
            "Metadata server issues short-lived tokens — protect the VM",
            "MIG = template + autoscaler + health checks + autohealing",
        ],
    )

    add_diagram_slide(
        prs,
        "Regional MIG + Global Load Balancer",
        [
            "        [ Users ]",
            "            ↓",
            "    Global HTTP(S) Load Balancer",
            "       ↙            ↘",
            "  [ VM zone-a ]   [ VM zone-b ]",
            "       Regional Managed Instance Group",
            "            ↑",
            "      Autoscaler (CPU 60%)",
        ],
        notes="Regional MIG survives zone failure. Single VM survives almost nothing.",
    )

    # --- Week 4 ---
    add_section_slide(prs, "Week 4", "Containers: GKE & Cloud Run")

    add_diagram_slide(
        prs,
        "Compute decision tree",
        [
            "Need full OS / GPU you manage?  →  GCE / MIG",
            "Many services, sidecars, operators?  →  GKE",
            "HTTP/events, scale to zero OK?  →  Cloud Run",
            "Tiny glue (GCS finalize → API)?  →  Cloud Functions",
            "",
            "Default for student projects: Cloud Run",
        ],
        notes="Start at Cloud Run. Move left only when a constraint appears.",
    )

    add_two_column_slide(
        prs,
        "GKE vs Cloud Run",
        "GKE (Kubernetes)",
        [
            "Regional control plane (Google-managed)",
            "Node pools on GCE VMs",
            "Workload Identity (no JSON keys)",
            "Best: complex networking, operators, mixed workloads",
        ],
        "Cloud Run",
        [
            "Run a container — HTTPS endpoint",
            "Scale to zero, concurrency per instance",
            "Revision-based deploys & canary splits",
            "Best: single HTTP microservices, bursty traffic",
        ],
    )

    # --- Week 5 ---
    add_section_slide(prs, "Week 5", "VPC Networking")

    add_bullet_slide(
        prs,
        "VPC model (GCP vs AWS trap)",
        [
            "GCP: VPC is GLOBAL; subnets are REGIONAL",
            "Custom-mode VPC required for production hygiene",
            "Firewall rules: stateful, by network tags or service accounts",
            "Private Google Access: reach *.googleapis.com without public IP",
            "Cloud NAT: outbound Internet for private VMs",
            "IAP: authenticate before SSH — never 0.0.0.0/0 on port 22",
            "Peering is NOT transitive (A–B and B–C ≠ A–C)",
        ],
    )

    add_diagram_slide(
        prs,
        "Private VM connectivity",
        [
            "GCE (no public IP)",
            "   ├─ Private Google Access → GCS / Pub/Sub APIs",
            "   ├─ Cloud NAT → Internet (apt, third-party)",
            "   └─ IAP tunnel → SSH from authorized users only",
        ],
    )

    # --- Week 6 ---
    add_section_slide(prs, "Week 6", "Load Balancing, CDN & Security Edge")

    add_table_slide(
        prs,
        "Load balancer taxonomy",
        ["Type", "Use", "GCP product"],
        [
            ["Application (L7)", "HTTP/HTTPS, SSL, URL maps", "Global HTTP(S) LB"],
            ["Passthrough (L4)", "TCP/UDP, preserve client IP", "Network LB"],
            ["Internal", "East-west RFC1918 traffic", "Internal LB"],
        ],
    )

    add_bullet_slide(
        prs,
        "CDN, DNS & Cloud Armor",
        [
            "Cloud CDN / Media CDN: cache at edge, signed URLs for premium content",
            "Cloud DNS: public/private zones, weighted & geo routing",
            "Cloud Armor: WAF + DDoS on global HTTP(S) LB",
            "Health checks must hit a cheap /healthz endpoint",
            "Prefer stateless backends over sticky sessions",
        ],
    )

    # --- Week 7 ---
    add_section_slide(prs, "Week 7", "Object & Block Storage")

    add_table_slide(
        prs,
        "Storage comparison",
        ["Product", "Semantics", "Best for"],
        [
            ["Cloud Storage (GCS)", "Object", "Data lake, backups, static web, signed uploads"],
            ["Persistent Disk", "Block", "VM boot, single-writer databases"],
            ["Filestore", "POSIX NFS", "Lift-and-shift shared files"],
        ],
    )

    add_bullet_slide(
        prs,
        "GCS essentials",
        [
            "Classes: Standard, Nearline, Coldline, Archive (+ Autoclass)",
            "Strong consistency for object overwrites",
            "Uniform bucket-level IAM (not public ACLs)",
            "Signed URLs: browser upload without public bucket",
            "Lifecycle rules: expire tmp/ after 7 days in labs",
            "Versioning + CMEK for regulated data",
        ],
    )

    # --- Week 8 ---
    add_section_slide(prs, "Week 8", "Databases — The Decision Lecture")

    add_diagram_slide(
        prs,
        "Database decision tree",
        [
            "SQL transactions, one region     →  Cloud SQL",
            "SQL global, horizontal scale     →  Spanner",
            "Documents + live client sync     →  Firestore",
            "Huge keyed telemetry / IoT       →  Bigtable",
            "Sub-ms cache / sessions          →  Memorystore",
            "Analytics SQL on TB–PB           →  BigQuery",
        ],
        notes="Do not put checkout transactions in BigQuery. Do not scan Firestore for analytics.",
    )

    add_bullet_slide(
        prs,
        "CAP & hot-spotting (exam favorites)",
        [
            "You cannot maximize C, A, and P independently",
            "Spanner: TrueTime waits out clock uncertainty for global commits",
            "Firestore: strong per-document; realtime listeners push to clients",
            "Bigtable: row key design is the course — avoid monotonic keys",
            "Hot geohash / hot user ID = one tablet overloaded",
        ],
    )

    # --- Week 9 ---
    add_section_slide(prs, "Week 9", "Events & Stream Processing")

    add_diagram_slide(
        prs,
        "Pub/Sub fan-out (real-time backbone)",
        [
            "  [ Producers: API / IoT / GCE ]",
            "              ↓",
            "         [ Topic ]",
            "        ↙   ↓   ↘",
            "  Cloud Run  Dataflow  BigQuery sub",
            "        ↘",
            "    Dead Letter Topic",
            "",
            "One topic, many independent subscribers",
        ],
        notes="Analytics must never block the checkout HTTP path.",
    )

    add_bullet_slide(
        prs,
        "Pub/Sub & Dataflow",
        [
            "At-least-once delivery — design idempotent consumers",
            "Ordering keys: per-key order, throughput trade-off",
            "Dead letter topics for poison messages",
            "Seek / replay for operations and debugging",
            "Dataflow (Apache Beam): batch + streaming, windows, watermarks",
            "Event time ≠ processing time — late data still happens",
        ],
    )

    add_diagram_slide(
        prs,
        "Idempotency pattern",
        [
            "Publisher → Topic → Consumer",
            "Consumer crashes before ACK",
            "Topic redelivers same message",
            "",
            "Fix: UPSERT with idempotency key",
            "  INSERT ... ON CONFLICT (order_id) DO NOTHING",
        ],
    )

    # --- Week 10 ---
    add_section_slide(prs, "Week 10", "Analytics & Machine Learning")

    add_bullet_slide(
        prs,
        "BigQuery",
        [
            "Storage (columnar) independent of compute (slots / on-demand)",
            "Standard SQL — partition + cluster or pay for full scans",
            "Streaming inserts vs load jobs vs Storage Write API",
            "BigQuery ML: CREATE MODEL for quick baselines",
            "Looker Studio for dashboards",
            "Teach students: LIMIT, preview, _TABLE_SUFFIX on public data",
        ],
    )

    add_table_slide(
        prs,
        "Vertex AI (survey)",
        ["Component", "Role"],
        [
            ["Training", "Produce model artifacts"],
            ["Feature Store", "Online/offline features for scoring"],
            ["Endpoints", "Low-latency online prediction"],
            ["Batch prediction", "Score a table overnight"],
            ["Vector Search / Gemini", "RAG, semantic retrieval, LLM apps"],
        ],
        notes="Real-time ML: event → features → Vertex endpoint → decision → log to BQ.",
    )

    # --- Week 11 ---
    add_section_slide(prs, "Week 11", "Security & Operations")

    add_bullet_slide(
        prs,
        "Defense in depth",
        [
            "Organization Policy: disable public buckets, restrict regions",
            "VPC Service Controls: anti-exfiltration perimeter",
            "Cloud KMS / CMEK / Cloud HSM",
            "Secret Manager — never secrets in Git or metadata",
            "Binary Authorization + Artifact vulnerability scanning",
            "Cloud Armor + IAP on public admin paths",
        ],
    )

    add_table_slide(
        prs,
        "Observability — four golden signals",
        ["Signal", "GCP tools"],
        [
            ["Latency", "Cloud Monitoring, Cloud Trace"],
            ["Traffic", "Cloud Monitoring request counts"],
            ["Errors", "Error Reporting, log-based metrics"],
            ["Saturation", "CPU/memory/disk metrics, slot usage"],
        ],
        notes="Propagate trace ID across HTTP and Pub/Sub or debugging is impossible.",
    )

    add_diagram_slide(
        prs,
        "CI/CD without JSON keys",
        [
            "GitHub Actions",
            "    ↓ OIDC token",
            "Workload Identity Federation",
            "    ↓ short-lived GCP token",
            "Artifact Registry → Cloud Run deploy",
        ],
    )

    # --- Week 12 ---
    add_section_slide(prs, "Week 12", "Architecture, Cost & Revision")

    add_bullet_slide(
        prs,
        "Well-architected questions",
        [
            "What fails? (zone, region, identity, quota, poison message)",
            "What are RPO and RTO?",
            "What consistency does the business require?",
            "What is the p99 latency budget and where is it spent?",
            "What is unit cost at 10× traffic?",
            "Who can exfiltrate data?",
        ],
    )

    add_diagram_slide(
        prs,
        "Student SaaS reference architecture",
        [
            "Browser → CDN → HTTPS LB + Armor → Cloud Run API",
            "   ├─ Cloud SQL (private IP)",
            "   ├─ GCS signed URLs",
            "   └─ Pub/Sub",
            "         ├─ Search indexer (Cloud Run)",
            "         └─ Dataflow → BigQuery → Looker Studio",
        ],
        notes="Capstone must justify every box. Extra services need a sentence why.",
    )

    add_table_slide(
        prs,
        "Cost landmines",
        ["Risk", "Mitigation"],
        [
            ["Idle GKE nodes", "Delete cluster after lab; use Cloud Run default"],
            ["BigQuery SELECT *", "Partition, cluster, column pruning"],
            ["Internet egress", "CDN, same-region compute+data"],
            ["Forgotten disks / static IPs", "Tear-down checklist every session"],
        ],
    )

    # --- Real-time use cases ---
    add_section_slide(prs, "Seminars", "Real-Time Use Cases")

    add_table_slide(
        prs,
        "Real-time case studies",
        ["Domain", "Sync path", "Async / stream"],
        [
            ["Location game", "GKE/Run + global LB", "Pub/Sub → Dataflow → BQ"],
            ["Ride matching", "Dispatch API", "Location stream join (Dataflow)"],
            ["Card fraud", "Vertex + Redis <40ms", "Score logs to BQ"],
            ["Campus IoT", "—", "Pub/Sub → Dataflow → alerts"],
            ["Live sports / OTT", "Media CDN", "Scores via Pub/Sub → Firestore"],
            ["Social media ingest", "Signed URL to GCS", "Object-finalize workers"],
        ],
        notes="Seminar Q: what is synchronous vs asynchronous in each case?",
    )

    add_two_column_slide(
        prs,
        "Case deep-dive: fraud scoring",
        "Requirements",
        [
            "p99 < 40 ms inside VPC",
            "Velocity counters (tx / 10 min)",
            "Model refresh without downtime",
            "Every score logged for audit",
        ],
        "GCP mapping",
        [
            "Cloud Run min instances or GKE",
            "Memorystore for velocity",
            "Vertex private endpoint",
            "Pub/Sub async enrichment (not on sync path)",
            "BigQuery training + drift monitoring",
        ],
    )

    add_two_column_slide(
        prs,
        "Case deep-dive: IoT campus",
        "Requirements",
        [
            "10k sensors every 5 seconds",
            "Alerts in < 15 seconds",
            "Years of history for analytics",
            "Device credential rotation",
        ],
        "GCP mapping",
        [
            "Pub/Sub ingest",
            "Dataflow rules + windows",
            "Bigtable (hot) + BigQuery (cold)",
            "Monitoring + Pub/Sub → notifier",
            "Firestore device twins",
        ],
    )

    # --- Labs & assessment ---
    add_section_slide(prs, "Labs", "Hands-On & Assessment")

    add_bullet_slide(
        prs,
        "Core lab: real-time spine",
        [
            "Dataset + partitioned table in BigQuery",
            "Pub/Sub topic + push/pull subscription",
            "Cloud Run consumer with idempotent insert",
            "Publish 100 messages — verify COUNT(*) in BQ",
            "Poison message → dead letter topic after N nacks",
            "Extension: Dataflow template Pub/Sub → BigQuery",
        ],
    )

    add_table_slide(
        prs,
        "Assessment (suggested)",
        ["Component", "Weight"],
        [
            ["Lab portfolio (8 labs)", "30%"],
            ["Midterm (weeks 1–6)", "20%"],
            ["Capstone + demo", "30%"],
            ["Final exam", "20%"],
        ],
    )

    add_bullet_slide(
        prs,
        "Capstone requirements",
        [
            "Must include real-time path (Pub/Sub, Dataflow, or Firestore listeners)",
            "IAM: no Editor on runtime service account",
            "Custom VPC or documented serverless egress story",
            "Failure modes: DLQ, retries, idempotency",
            "Cost estimate at 10× traffic",
            "8-minute demo + tear-down runbook",
        ],
    )

    # --- Closing ---
    add_bullet_slide(
        prs,
        "Key takeaways",
        [
            "Cloud = APIs + elasticity + shared responsibility",
            "Default compute: Cloud Run unless GCE/GKE constraints appear",
            "One Pub/Sub topic, many subscribers — decouple sync from async",
            "At-least-once → idempotent consumers → dead letter queues",
            "BigQuery for analytics; Cloud SQL/Spanner for transactions",
            "No JSON keys; budget alerts; delete GKE after every lab",
        ],
        notes="Thank students. Point to gcp-bachelors-course/ for full notes, figures, labs.",
    )

    add_title_slide(
        prs,
        "Thank you",
        "Full materials: gcp-bachelors-course/\nLecture notes · Mermaid figures · Labs · Assessments",
    )

    prs.save(OUT)
    print(f"Wrote {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
